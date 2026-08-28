"""Native, structure-preserving parsers for non-PDF Office documents."""

from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from io import BytesIO
from typing import TYPE_CHECKING, Any

from .elements import Element

if TYPE_CHECKING:
    from docx.document import Document as DocxDocument


@dataclass(frozen=True)
class NativeDocumentArtifact:
    parser_name: str
    parser_version: str
    source_sha256: str
    elements: list[Element]
    assets: list[dict[str, object]] = field(default_factory=list)


def parse_docx_bytes(content: bytes, *, document_id: str, source_url: str = "") -> NativeDocumentArtifact:
    """Map DOCX blocks into typed elements without inventing page coordinates."""

    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(BytesIO(content))
    source_sha256 = hashlib.sha256(content).hexdigest()
    parser_version = str(docx.__version__)
    blocks = list(document.iter_inner_content())
    elements: list[Element] = []
    assets: list[dict[str, object]] = []
    heading_path: list[str] = []
    order = 0

    def emit(
        element_type: str,
        text: str,
        *,
        sub_type: str = "",
        image_path: str = "",
        caption: str = "",
    ) -> None:
        nonlocal order
        elements.append(
            Element(
                document_id=document_id,
                element_id=f"{document_id}:docx:e{order}",
                reading_order=order,
                type=element_type,  # type: ignore[arg-type]
                sub_type=sub_type,
                heading_path=list(heading_path),
                page_index=0,
                bbox=[0.0, 0.0, 0.0, 0.0],
                coordinate_system="document",
                text=text,
                image_path=image_path,
                caption=caption,
                parser_name="python-docx",
                parser_version=parser_version,
                source_sha256=source_sha256,
                source_url=source_url,
            )
        )
        order += 1

    index = 0
    while index < len(blocks):
        block = blocks[index]
        if isinstance(block, Table):
            rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in block.rows]
            if any(row.strip(" |") for row in rows):
                emit("table", "\n".join(rows), sub_type="table")
            index += 1
            continue
        if not isinstance(block, Paragraph):
            index += 1
            continue

        style_name = (block.style.name or "").strip()
        text = block.text.strip()
        image_refs = _docx_image_refs(document, block)
        if image_refs:
            caption = ""
            if index + 1 < len(blocks) and isinstance(blocks[index + 1], Paragraph):
                following = blocks[index + 1]
                if (following.style.name or "").strip().lower() == "caption":
                    caption = following.text.strip()
                    index += 1
            for image_ref in image_refs:
                emit("image", text, sub_type="inline_image", image_path=image_ref, caption=caption)
                assets.append({"element_id": elements[-1].element_id, "path": image_ref})
            index += 1
            continue
        if not text:
            index += 1
            continue
        if style_name.lower().startswith("heading"):
            level = _heading_level(style_name)
            heading_path = heading_path[: level - 1] + [text]
            emit("heading", text, sub_type=f"heading_{level}")
        elif style_name.lower().startswith("list"):
            emit("list", text, sub_type=style_name)
        elif style_name.lower() != "caption":
            emit("text", text, sub_type="paragraph")
        index += 1

    return NativeDocumentArtifact("python-docx", parser_version, source_sha256, elements, assets)


def parse_pptx_bytes(content: bytes, *, document_id: str, source_url: str = "") -> NativeDocumentArtifact:
    """Map PPTX slides and shapes into typed elements with slide-relative coordinates."""

    import pptx

    presentation = pptx.Presentation(BytesIO(content))
    source_sha256 = hashlib.sha256(content).hexdigest()
    parser_version = str(pptx.__version__)
    elements: list[Element] = []
    assets: list[dict[str, object]] = []
    order = 0

    def emit(
        element_type: str,
        text: str,
        *,
        slide_index: int,
        heading_path: list[str],
        sub_type: str = "",
        bbox: list[float] | None = None,
        image_path: str = "",
        caption: str = "",
    ) -> None:
        nonlocal order
        elements.append(
            Element(
                document_id=document_id,
                element_id=f"{document_id}:slide:{slide_index}:e{order}",
                reading_order=order,
                type=element_type,  # type: ignore[arg-type]
                sub_type=sub_type,
                heading_path=list(heading_path),
                page_index=slide_index,
                bbox=bbox or [0.0, 0.0, 0.0, 0.0],
                coordinate_system="slide_1000",
                text=text,
                image_path=image_path,
                caption=caption,
                parser_name="python-pptx",
                parser_version=parser_version,
                source_sha256=source_sha256,
                source_url=source_url,
            )
        )
        order += 1

    for slide_index, slide in enumerate(presentation.slides):
        title_text = (slide.shapes.title.text if slide.shapes.title is not None else "").strip()
        slide_title = title_text or f"Slide {slide_index + 1}"
        heading_path = [slide_title]
        emit("heading", slide_title, slide_index=slide_index, heading_path=heading_path, sub_type="slide")
        for shape in _iter_pptx_shapes(slide.shapes):
            bbox = _slide_bbox(shape, presentation.slide_width, presentation.slide_height)
            if getattr(shape, "has_table", False):
                rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in shape.table.rows]
                emit("table", "\n".join(rows), slide_index=slide_index, heading_path=heading_path, sub_type="table", bbox=bbox)
                continue
            if getattr(shape, "has_chart", False):
                emit("image", "", slide_index=slide_index, heading_path=heading_path, sub_type="chart", bbox=bbox, caption="Chart")
                continue
            image_path = _pptx_image_ref(shape)
            if image_path:
                emit("image", "", slide_index=slide_index, heading_path=heading_path, sub_type="image", bbox=bbox, image_path=image_path)
                assets.append({"element_id": elements[-1].element_id, "path": image_path})
                continue
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    subtype = "title" if shape is slide.shapes.title else "text"
                    emit("text", text, slide_index=slide_index, heading_path=heading_path, sub_type=subtype, bbox=bbox)
        notes = (slide.notes_slide.notes_text_frame.text or "").strip()
        if notes:
            emit("text", notes, slide_index=slide_index, heading_path=heading_path, sub_type="speaker_notes")

    return NativeDocumentArtifact("python-pptx", parser_version, source_sha256, elements, assets)


def _heading_level(style_name: str) -> int:
    try:
        return max(1, int(style_name.rsplit(" ", 1)[-1]))
    except ValueError:
        return 1


def _docx_image_refs(document: DocxDocument, paragraph: Any) -> list[str]:
    refs: list[str] = []
    for blip in paragraph._p.xpath(".//a:blip"):
        relation_id = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
        if not relation_id:
            continue
        part = document.part.related_parts.get(relation_id)
        blob = getattr(part, "blob", None)
        if isinstance(blob, bytes):
            refs.append(f"embedded://sha256:{hashlib.sha256(blob).hexdigest()}")
    return list(dict.fromkeys(refs))


def _iter_pptx_shapes(shapes: Any):
    for shape in shapes:
        if getattr(shape, "shape_type", None) is not None and getattr(shape, "shapes", None) is not None:
            yield from _iter_pptx_shapes(shape.shapes)
        else:
            yield shape


def _pptx_image_ref(shape: Any) -> str:
    try:
        blob = shape.image.blob
    except (AttributeError, ValueError):
        return ""
    return f"embedded://sha256:{hashlib.sha256(blob).hexdigest()}" if blob else ""


def _slide_bbox(shape: Any, width: int, height: int) -> list[float]:
    if width <= 0 or height <= 0:
        return [0.0, 0.0, 0.0, 0.0]
    left = max(0.0, min(1000.0, float(getattr(shape, "left", 0)) * 1000 / width))
    top = max(0.0, min(1000.0, float(getattr(shape, "top", 0)) * 1000 / height))
    right = max(left, min(1000.0, float(getattr(shape, "left", 0) + getattr(shape, "width", 0)) * 1000 / width))
    bottom = max(top, min(1000.0, float(getattr(shape, "top", 0) + getattr(shape, "height", 0)) * 1000 / height))
    return [left, top, right, bottom]
