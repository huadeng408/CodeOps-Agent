from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from io import BytesIO

import openpyxl
from openpyxl.utils import get_column_letter

from .elements import Element


@dataclass(frozen=True)
class SpreadsheetRegion:
    sheet_name: str
    range_ref: str
    row_count: int
    column_count: int


@dataclass(frozen=True)
class SpreadsheetArtifact:
    parser_name: str
    parser_version: str
    source_sha256: str
    elements: list[Element]
    formula_dependencies: dict[str, tuple[str, ...]] = field(default_factory=dict)
    named_ranges: dict[str, tuple[str, ...]] = field(default_factory=dict)
    cached_values: dict[str, object] = field(default_factory=dict)
    formula_cache_present: dict[str, bool] = field(default_factory=dict)
    table_regions: tuple[SpreadsheetRegion, ...] = ()


def parse_xlsx_bytes(content: bytes, *, document_id: str, source_url: str = "") -> SpreadsheetArtifact:
    """Extract bounded regions, formulas, names and cache state without recalculation."""

    source_sha256 = hashlib.sha256(content).hexdigest()
    parser_version = openpyxl.__version__
    workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False)
    cached_workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    elements: list[Element] = []
    formula_dependencies: dict[str, tuple[str, ...]] = {}
    cached_values: dict[str, object] = {}
    formula_cache_present: dict[str, bool] = {}
    regions: list[SpreadsheetRegion] = []
    order = 0

    for sheet in workbook.worksheets:
        cached_sheet = cached_workbook[sheet.title]
        heading_path = [sheet.title]
        elements.append(_element(document_id, order, "heading", sheet.title, heading_path, source_sha256, parser_version, source_url, sheet_name=sheet.title))
        order += 1
        rows = list(sheet.iter_rows())
        max_row = len(rows)
        max_col = max((len(row) for row in rows), default=0)
        formula_records: list[tuple[str, str, str, tuple[str, ...], object]] = []
        for row in rows:
            for cell in row:
                if not isinstance(cell.value, str) or not cell.value.startswith("="):
                    continue
                ref = f"{sheet.title}!{cell.coordinate}"
                cached_cell = cached_sheet[cell.coordinate]
                cached_values[ref] = cached_cell.value
                formula_cache_present[ref] = cached_cell.value is not None
                formula_records.append((ref, cell.value, cell.coordinate, _formula_dependencies(cell.value), cached_cell.value))
        if not rows:
            continue
        for start_row in range(1, max_row + 1, 40):
            end_row = min(max_row, start_row + 39)
            end_col = min(max_col, 12)
            range_ref = f"{sheet.title}!A{start_row}:{get_column_letter(end_col)}{end_row}"
            region = SpreadsheetRegion(sheet.title, range_ref, end_row - start_row + 1, end_col)
            regions.append(region)
            lines = []
            for row_number in range(start_row, end_row + 1):
                values = [_cell_text(sheet.cell(row=row_number, column=column).value) for column in range(1, end_col + 1)]
                lines.append(" | ".join(values))
            elements.append(_element(
                document_id, order, "table",
                    f"Sheet: {sheet.title}\nRange: {range_ref.split('!', 1)[1]} ({range_ref})\n" + "\n".join(lines),
                heading_path, source_sha256, parser_version, source_url,
                sheet_name=sheet.title, cell_range=range_ref,
            ))
            order += 1
        for ref, formula, coordinate, dependencies, _cached in formula_records:
            elements.append(_element(
                document_id, order, "equation", f"{ref} {formula}", heading_path,
                source_sha256, parser_version, source_url, latex=formula,
                sheet_name=sheet.title, cell_range=coordinate,
            ))
            formula_dependencies[ref] = dependencies
            order += 1

    named_ranges: dict[str, tuple[str, ...]] = {}
    for name, defined in workbook.defined_names.items():
        destinations = tuple(f"{sheet_name}!{cell_range}" for sheet_name, cell_range in defined.destinations)
        if destinations:
            named_ranges[str(name)] = destinations
    return SpreadsheetArtifact(
        "openpyxl", parser_version, source_sha256, elements, formula_dependencies,
        named_ranges, cached_values, formula_cache_present, tuple(regions),
    )


def _element(
    document_id: str,
    order: int,
    element_type: str,
    text: str,
    heading_path: list[str],
    source_sha256: str,
    parser_version: str,
    source_url: str,
    *,
    latex: str = "",
    sheet_name: str = "",
    cell_range: str = "",
) -> Element:
    return Element(
        document_id=document_id,
        element_id=f"{document_id}:sheet:{heading_path[0]}:e{order}",
        reading_order=order,
        type=element_type,  # type: ignore[arg-type]
        heading_path=heading_path,
        text=text,
        latex=latex,
        parser_name="openpyxl",
        parser_version=parser_version,
        source_sha256=source_sha256,
        source_url=source_url,
        sheet_name=sheet_name,
        cell_range=cell_range,
    )


def _cell_text(value: object) -> str:
    return "" if value is None else str(value)


def _formula_dependencies(formula: str) -> tuple[str, ...]:
    references = re.findall(r"(?:'[^']+'|[A-Za-z_][\w ]*)?!?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?", formula.upper())
    return tuple(dict.fromkeys(reference.replace("$", "") for reference in references))
