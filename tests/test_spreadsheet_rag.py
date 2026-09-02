from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace

from openpyxl import Workbook
import pytest

from orchestrator.rag.ingestion import IngestionService
from orchestrator.rag.models import ParseRequestPayload
from orchestrator.rag.spreadsheet import parse_xlsx_bytes


def test_xlsx_parser_preserves_sheet_table_and_formula_coordinates() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Budget"
    sheet.append(["Item", "Amount", "Double"])
    sheet.append(["Servers", 10, "=B2*2"])
    sheet.append(["Storage", 20, "=SUM(B2:B3)"])
    sheet.merge_cells("A5:C5")
    sheet["A5"] = "Quarterly plan"
    payload = BytesIO()
    workbook.save(payload)

    artifact = parse_xlsx_bytes(payload.getvalue(), document_id="budget.xlsx")

    assert artifact.parser_name == "openpyxl"
    assert artifact.source_sha256
    assert [element.type for element in artifact.elements] == ["heading", "table", "equation", "equation"]
    table = artifact.elements[1]
    assert "Sheet: Budget" in table.text
    assert "Range: A1:C5" in table.text
    assert "Item | Amount | Double" in table.text
    assert table.parser_version == artifact.parser_version
    assert table.source_sha256 == artifact.source_sha256
    assert artifact.elements[2].latex == "=B2*2"
    assert artifact.elements[2].text == "Budget!C2 =B2*2"
    assert artifact.elements[3].latex == "=SUM(B2:B3)"
    assert artifact.formula_dependencies == {
        "Budget!C2": ("B2",),
        "Budget!C3": ("B2:B3",),
    }


def test_xlsx_parser_preserves_named_ranges_cache_state_and_bounded_regions() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Metrics"
    sheet.append(["Metric", "Value", "Delta"])
    for index in range(1, 46):
        sheet.append([f"m{index}", index, f"=B{index + 1}*2"])
    sheet["B2"] = 10
    workbook.create_named_range("important_metrics", sheet, "A2:C5")
    payload = BytesIO()
    workbook.save(payload)

    artifact = parse_xlsx_bytes(payload.getvalue(), document_id="metrics.xlsx")

    assert artifact.named_ranges == {"important_metrics": ("Metrics!A2:C5",)}
    assert "Metrics!C2" in artifact.cached_values
    assert artifact.formula_cache_present["Metrics!C2"] is False
    assert len(artifact.table_regions) >= 2
    assert all(region.row_count <= 40 and region.column_count <= 12 for region in artifact.table_regions)
    assert artifact.table_regions[0].range_ref == "Metrics!A1:C40"
    assert artifact.table_regions[1].range_ref == "Metrics!A41:C46"
    table_elements = [element for element in artifact.elements if element.type == "table"]
    assert len(table_elements) == len(artifact.table_regions)
    assert all(region.range_ref in element.text for region, element in zip(artifact.table_regions, table_elements))


class _Response:
    def __init__(self, content: bytes) -> None:
        self.content = content
        self.text = ""

    def raise_for_status(self) -> None:
        return None


class _XlsxHTTP:
    def __init__(self, content: bytes) -> None:
        self.content = content
        self.tika_calls = 0

    async def get(self, _url: str) -> _Response:
        return _Response(self.content)

    async def put(self, *_args: object, **_kwargs: object) -> _Response:
        self.tika_calls += 1
        raise AssertionError("XLSX must use native openpyxl extraction")


@pytest.mark.asyncio
async def test_xlsx_parse_uses_native_extractor_and_never_tika() -> None:
    workbook = Workbook()
    workbook.active.title = "Budget"
    workbook.active.append(["Amount"])
    workbook.active.append(["=1+1"])
    payload = BytesIO()
    workbook.save(payload)
    http = _XlsxHTTP(payload.getvalue())
    service = IngestionService.__new__(IngestionService)
    service._http = http
    service._settings = SimpleNamespace(tika_url="http://tika.invalid")

    response = await service.parse(
        ParseRequestPayload.model_validate(
            {
                "task": {"file_md5": "budget", "file_name": "budget.xlsx", "user_id": 1, "stage": "parse"},
                "objectUrl": "http://source.invalid/budget.xlsx",
            }
        )
    )

    assert http.tika_calls == 0
    assert response.parserName == "openpyxl"
    assert [element.type for element in response.elements] == ["heading", "table", "equation"]
    assert response.tableRegions[0]["range_ref"] == "Budget!A1:A2"
    assert response.formulaCachePresent["Budget!A2"] is False


@pytest.mark.asyncio
async def test_xlsx_parse_preserves_stable_corpus_document_id() -> None:
    workbook = Workbook()
    workbook.active.title = "Budget"
    workbook.active.append(["Amount"])
    payload = BytesIO()
    workbook.save(payload)
    service = IngestionService.__new__(IngestionService)
    service._http = _XlsxHTTP(payload.getvalue())
    service._settings = SimpleNamespace(tika_url="http://tika.invalid")
    document_id = "go@0123456789abcdef0123456789abcdef01234567:docs/budget.xlsx"

    response = await service.parse(
        ParseRequestPayload.model_validate(
            {
                "task": {
                    "file_md5": "file-md5-not-document-id",
                    "document_id": document_id,
                    "file_name": "budget.xlsx",
                    "user_id": 1,
                    "stage": "parse",
                },
                "objectUrl": "http://source.invalid/budget.xlsx",
            }
        )
    )

    assert response.documentId == document_id
    assert {element.document_id for element in response.elements} == {document_id}
