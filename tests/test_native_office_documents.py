from __future__ import annotations

import base64
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document as DocxDocument
from docx.shared import Inches
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches as PptxInches

from orchestrator.rag.ingestion import IngestionService
from orchestrator.rag.models import ParseRequestPayload
from orchestrator.rag.native_documents import parse_docx_bytes, parse_pptx_bytes
from orchestrator.rag.models import SearchResultPayload
from orchestrator.rag.retrievers import search_result_to_document

FIXTURES = Path(__file__).parent / "fixtures" / "multimodal"


def _png_bytes() -> bytes:
    return base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/"
        "e+m+7wAAAABJRU5ErkJggg=="
    )


def _docx_bytes() -> bytes:
    document = DocxDocument()
    document.add_heading("Operations", level=1)
    document.add_paragraph("Deploy the service from the release package.")
    document.add_paragraph("Verify health before traffic.", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Check"
    table.cell(0, 1).text = "Command"
    table.cell(1, 0).text = "Health"
    table.cell(1, 1).text = "curl /health"
    document.add_picture(BytesIO(_png_bytes()), width=Inches(0.2))
    document.add_paragraph("Deployment diagram", style="Caption")
    return _save_docx(document)


def _save_docx(document: DocxDocument) -> bytes:
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _pptx_bytes() -> bytes:
    presentation = Presentation()
    presentation.slide_width = PptxInches(10)
    presentation.slide_height = PptxInches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    title = slide.shapes.title
    title.text = "Release plan"
    body = slide.shapes.add_textbox(PptxInches(1), PptxInches(1.5), PptxInches(3), PptxInches(1))
    body.text_frame.text = "Deploy, verify, observe"
    table = slide.shapes.add_table(2, 2, PptxInches(1), PptxInches(3), PptxInches(4), PptxInches(1.5))
    table.table.cell(0, 0).text = "Stage"
    table.table.cell(0, 1).text = "Owner"
    table.table.cell(1, 0).text = "Verify"
    table.table.cell(1, 1).text = "SRE"
    slide.shapes.add_picture(BytesIO(_png_bytes()), PptxInches(6), PptxInches(1), width=PptxInches(1), height=PptxInches(1))
    slide.notes_slide.notes_text_frame.text = "Speaker note: mention rollback."
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


class _Response:
    def __init__(self, content: bytes = b"", text: str = "") -> None:
        self.content = content
        self.text = text

    def raise_for_status(self) -> None:
        return None


class _NoTikaHTTP:
    def __init__(self, source: bytes) -> None:
        self.source = source
        self.put_calls = 0

    async def get(self, _url: str) -> _Response:
        return _Response(content=self.source)

    async def put(self, *_args: object, **_kwargs: object) -> _Response:
        self.put_calls += 1
        raise AssertionError("native Office input must not call Tika")


def test_docx_parser_emits_stable_structured_elements() -> None:
    source = _docx_bytes()
    artifact = parse_docx_bytes(source, document_id="docx-1")
    repeated = parse_docx_bytes(source, document_id="docx-1")

    assert artifact.parser_name == "python-docx"
    assert artifact.parser_version
    assert artifact.source_sha256
    assert [item.type for item in artifact.elements] == ["heading", "text", "list", "table", "image"]
    assert [item.element_id for item in artifact.elements] == [item.element_id for item in repeated.elements]
    assert artifact.elements[0].heading_path == ["Operations"]
    assert "Check | Command" in artifact.elements[3].text
    assert artifact.elements[4].caption == "Deployment diagram"
    assert artifact.elements[4].image_path.startswith("embedded://sha256:")


def test_checked_in_office_fixtures_parse_with_stable_source_identity() -> None:
    docx_source = (FIXTURES / "office_chunk_fixture.docx").read_bytes()
    pptx_source = (FIXTURES / "office_chunk_fixture.pptx").read_bytes()
    docx = parse_docx_bytes(docx_source, document_id="fixture-docx")
    pptx = parse_pptx_bytes(pptx_source, document_id="fixture-pptx")

    assert docx.source_sha256 == parse_docx_bytes(docx_source, document_id="fixture-docx").source_sha256
    assert pptx.source_sha256 == parse_pptx_bytes(pptx_source, document_id="fixture-pptx").source_sha256
    assert {item.type for item in docx.elements} >= {"heading", "text", "list", "table", "image"}
    assert {item.sub_type for item in pptx.elements} >= {"slide", "text", "table", "image", "speaker_notes"}


def test_pptx_parser_emits_slide_shapes_notes_and_normalized_coordinates() -> None:
    source = _pptx_bytes()
    artifact = parse_pptx_bytes(source, document_id="pptx-1")
    repeated = parse_pptx_bytes(source, document_id="pptx-1")

    assert artifact.parser_name == "python-pptx"
    assert artifact.parser_version
    assert [item.element_id for item in artifact.elements] == [item.element_id for item in repeated.elements]
    assert artifact.elements[0].type == "heading"
    assert artifact.elements[0].page_index == 0
    assert {item.sub_type for item in artifact.elements} >= {"text", "table", "image", "speaker_notes"}
    shape_elements = artifact.elements[1:]
    assert all(0 <= value <= 1000 for item in shape_elements for value in item.bbox)
    assert any(item.type == "image" and item.image_path.startswith("embedded://sha256:") for item in shape_elements)
    assert any(item.sub_type == "speaker_notes" and "rollback" in item.text for item in shape_elements)


def test_search_result_preserves_spreadsheet_coordinates_for_retrieval() -> None:
    result = SearchResultPayload.model_validate(
        {
            "fileMd5": "metrics-md5",
            "fileName": "metrics.xlsx",
            "chunkId": 4,
            "textContent": "Metrics!A1:C40",
            "sheetName": "Metrics",
            "cellRange": "Metrics!A1:C40",
        }
    )
    document = search_result_to_document(result, "bm25")
    assert document.metadata["sheetName"] == "Metrics"
    assert document.metadata["cellRange"] == "Metrics!A1:C40"


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("file_name", "source_factory", "parser_name"),
    [("runbook.docx", _docx_bytes, "python-docx"), ("release.pptx", _pptx_bytes, "python-pptx")],
)
async def test_native_office_parse_route_never_calls_tika(
    file_name: str, source_factory, parser_name: str
) -> None:
    source = source_factory()
    http = _NoTikaHTTP(source)
    service = IngestionService.__new__(IngestionService)
    service._http = http
    service._settings = SimpleNamespace(tika_url="http://tika.invalid")
    result = await service.parse(
        ParseRequestPayload.model_validate(
            {
                "task": {"file_md5": "office-1", "file_name": file_name, "user_id": 1, "stage": "parse"},
                "objectUrl": f"http://objects.invalid/{file_name}",
            }
        )
    )

    assert http.put_calls == 0
    assert result.parserName == parser_name
    assert result.documentId == "office-1"
    assert result.elements
