"""Generate small deterministic-shape Office fixtures for parser integration tests."""

from __future__ import annotations

import base64
from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from docx.shared import Inches
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches as PptxInches

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/"
    "e+m+7wAAAABJRU5ErkJggg=="
)


def save_docx(path: Path) -> None:
    document = DocxDocument()
    document.add_heading("Operations", level=1)
    document.add_paragraph("Deploy the service from the release package.")
    document.add_paragraph("Verify health before traffic.", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Check"
    table.cell(0, 1).text = "Command"
    table.cell(1, 0).text = "Health"
    table.cell(1, 1).text = "curl /health"
    document.add_picture(BytesIO(PNG), width=Inches(0.2))
    document.add_paragraph("Deployment diagram", style="Caption")
    document.save(path)


def save_pptx(path: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = PptxInches(10)
    presentation.slide_height = PptxInches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Release plan"
    body = slide.shapes.add_textbox(PptxInches(1), PptxInches(1.5), PptxInches(3), PptxInches(1))
    body.text_frame.text = "Deploy, verify, observe"
    table = slide.shapes.add_table(2, 2, PptxInches(1), PptxInches(3), PptxInches(4), PptxInches(1.5))
    table.table.cell(0, 0).text = "Stage"
    table.table.cell(0, 1).text = "Owner"
    table.table.cell(1, 0).text = "Verify"
    table.table.cell(1, 1).text = "SRE"
    slide.shapes.add_picture(BytesIO(PNG), PptxInches(6), PptxInches(1), width=PptxInches(1), height=PptxInches(1))
    slide.notes_slide.notes_text_frame.text = "Speaker note: mention rollback."
    presentation.save(path)


def save_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Metrics"
    sheet.append(["Metric", "Value", "Delta"])
    for index in range(1, 46):
        sheet.append([f"m{index}", index, f"=B{index + 1}*2"])
    workbook.create_named_range("important_metrics", sheet, "A2:C5")
    workbook.save(path)


def main() -> None:
    target = Path(__file__).parents[2] / "tests" / "fixtures" / "multimodal"
    target.mkdir(parents=True, exist_ok=True)
    save_docx(target / "office_chunk_fixture.docx")
    save_pptx(target / "office_chunk_fixture.pptx")
    save_xlsx(target / "office_chunk_fixture.xlsx")


if __name__ == "__main__":
    main()
